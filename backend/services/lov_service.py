"""Discovers LOV CSV files and resolves the LOV requirements used to build DQ8 rule configuration."""

import csv
import io
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from core.storage_layout import get_lov_upload_dir
from core.config import LOV_FILE_EXTENSION


# Uploaded files are stored with a UUID prefix so two people uploading
# "lov.csv" don't collide (see UploadFileSaver._sanitize_filename). That prefix
# is storage bookkeeping, so the UI shows the name the user actually picked.
_UPLOAD_PREFIX = re.compile(r"^[0-9a-f]{32}_")


def strip_upload_prefix(filename: str) -> str:
    """Return an uploaded filename without its storage UUID prefix."""
    return _UPLOAD_PREFIX.sub("", filename)


class LovService:
    """Handles LOV discovery and extraction of LOV requirements from profile maps."""

    @staticmethod
    def _extract_sheet_lovs(
        profile_map: Path,
        sheet: str,
    ) -> set[str]:
        """Extract DQ8 LOV names from a single worksheet."""

        workbook = load_workbook(profile_map, read_only=True, data_only=True)

        try:
            worksheet = workbook[sheet]
            rows = worksheet.iter_rows(values_only=True)
            headers = [str(column or "").strip() for column in next(rows, ())]

            rules_column_single = "Applicable Rule\n(Single ID)"
            rules_column_legacy = "Applicable Rules\n(comma-sep IDs)"
            parameters_column = "Rule Parameters\n(see Instructions)"

            rules_column = rules_column_single if rules_column_single in headers else rules_column_legacy

            if rules_column not in headers:
                return set()

            lovs: set[str] = set()
            for row in rows:
                record = dict(zip(headers, row))

                rules = str(record.get(rules_column) or "").strip().upper()
                rule_list = [r.strip() for r in rules.replace(",", "|").replace(";", "|").split("|") if r.strip()]

                if "DQ8" not in rule_list:
                    continue

                parameters = str(record.get(parameters_column) or "").strip()
                parts = [
                    part.strip()
                    for part in parameters.split("|")
                    if part.strip()
                ]
                if parts:
                    lovs.add(parts[0])

            return lovs
        finally:
            workbook.close()

    def extract_required_lovs(
        self,
        profile_map_path: str | Path | None,
    ) -> set[str]:
        """Extract LOV names required by DQ8 rules in a profile map."""

        required: set[str] = set()

        if not profile_map_path:
            return required

        profile_map = Path(profile_map_path)

        if not profile_map.exists():
            return required

        try:
            excel_file = load_workbook(
                profile_map,
                read_only=True,
                data_only=True,
            )

            try:
                for sheet in excel_file.sheetnames:
                    if sheet.lower() == "instructions":
                        continue

                    required.update(
                        self._extract_sheet_lovs(profile_map, sheet)
                    )
            finally:
                excel_file.close()

        except Exception as exc:
            print(
                "[WARN] Could not extract DQ8 LOVs "
                f"from profile map: {exc}"
            )

        return required

    # ── CSV parsing ──────────────────────────────────────────────

    @staticmethod
    def summarize_csv(text: str) -> list[dict[str, Any]]:
        """Summarize a LOV CSV as one entry per column.

        A LOV CSV is *wide*: each column is an independent list of values named
        by its header, so a single file can carry every LOV a job needs - one
        column per DQ8-checked column, however many tables they span. Columns
        are ragged by design; a shorter list just leaves trailing blanks.

        Each entry is {"name", "values_count"} for one column, in file order.
        """
        reader = csv.reader(io.StringIO(text))

        header = next(reader, None)
        if not header:
            return []

        names = [str(cell or "").strip() for cell in header]
        counts = [0] * len(names)

        for row in reader:
            for index in range(min(len(row), len(names))):
                if str(row[index] or "").strip():
                    counts[index] += 1

        return [
            {"name": name, "values_count": counts[index]}
            for index, name in enumerate(names)
            if name
        ]

    @staticmethod
    def _read_text(path: Path) -> str:
        """Read a LOV CSV from disk, tolerating a UTF-8 BOM."""
        return path.read_text(encoding="utf-8-sig")

    # ── Listing (storage aware) ──────────────────────────────────

    def list_lov_files(self) -> list[dict[str, Any]]:
        """List every uploaded LOV file and the LOV names it defines.

        Each entry: { file, display_name, lovs: [{name, values_count}] }.
        """
        return self._list_lov_files_local()

    def _list_lov_files_local(self) -> list[dict[str, Any]]:
        """List LOV files from the local uploads directory."""
        directory = get_lov_upload_dir()

        if not directory.is_dir():
            return []

        files: list[dict[str, Any]] = []

        for path in sorted(directory.iterdir()):
            if not path.is_file() or path.suffix.lower() != LOV_FILE_EXTENSION:
                continue

            try:
                lovs = self.summarize_csv(self._read_text(path))
            except Exception as exc:
                print(f"[WARN] Could not read LOV file {path.name}: {exc}")
                lovs = []

            files.append(self._build_entry(path.name, lovs))

        return files

    @staticmethod
    def _build_entry(filename: str, lovs: list[dict[str, Any]]) -> dict[str, Any]:
        """Shape one listing row for the API."""
        return {
            "file": filename,
            "display_name": strip_upload_prefix(filename),
            "lovs": lovs,
        }

    def get_lov_file(self, filename: str) -> dict[str, Any] | None:
        """Return one LOV file's listing entry, or None if it is not stored.

        Lets a caller that already knows the filename - the job wizard, right
        after uploading it - read back the LOV names inside. It reads that one
        file rather than filtering a full listing.
        """
        if not filename or "/" in filename or "\\" in filename or ".." in filename:
            return None

        if not filename.lower().endswith(LOV_FILE_EXTENSION):
            return None

        directory = get_lov_upload_dir()
        target = directory / filename

        if not target.is_file():
            return None

        try:
            lovs = self.summarize_csv(self._read_text(target))
        except Exception as exc:
            print(f"[WARN] Could not read LOV file {filename}: {exc}")
            lovs = []

        return self._build_entry(filename, lovs)

    def delete_lov_file(self, filename: str) -> bool:
        """Delete one LOV file from the uploads directory. False if it is absent."""
        directory = get_lov_upload_dir()
        target = directory / filename

        # Reject anything that escapes the upload directory before touching disk.
        if target.resolve().parent != directory.resolve():
            raise ValueError("Invalid filename")

        if not target.is_file():
            return False

        target.unlink()
        return True

    # ── Discovery for the pipeline ───────────────────────────────

    def discover_lov_files(
        self,
        lov_dir: str | Path | None = None,
        lov_file: str | None = None,
    ) -> dict[str, str]:
        """Map every LOV name found on disk to the CSV file that defines it.

        Every column of a LOV CSV is its own LOV, so one file can satisfy every
        DQ8 rule in a job. LovProvider picks the matching column out of the file
        by name when it reads it.

        `lov_file` narrows the scan to the single file a job was created with;
        without it every CSV in the directory is read, which is what jobs
        created before LOVs were per-job still expect.
        """

        discovered: dict[str, str] = {}

        directory = (
            Path(lov_dir)
            if lov_dir
            else get_lov_upload_dir()
        )

        if not directory.is_dir():
            return discovered

        try:
            for path in sorted(directory.iterdir()):
                if not path.is_file():
                    continue

                if path.suffix.lower() != LOV_FILE_EXTENSION:
                    continue

                if lov_file and path.name != lov_file:
                    continue

                try:
                    columns = self.summarize_csv(self._read_text(path))
                except Exception as exc:
                    print(
                        "[WARN] Could not read header from "
                        f"{path.name}, using filename: {exc}"
                    )
                    discovered.setdefault(
                        strip_upload_prefix(path.stem),
                        str(path),
                    )
                    continue

                for column in columns:
                    name = column["name"]

                    if name in discovered and discovered[name] != str(path):
                        print(
                            f"[WARN] LOV '{name}' is defined in more than one "
                            f"file; keeping '{Path(discovered[name]).name}'"
                        )
                        continue

                    discovered[name] = str(path)

        except Exception as exc:
            print(
                f"[WARN] LOV discovery failed: {exc}"
            )

        return discovered

    def build_lov_configuration(
        self,
        profile_map_path: str | Path | None,
        lov_dir: str | Path | None = None,
        lov_file: str | None = None,
    ) -> dict[str, str]:
        """Build the LOV configuration required by the DQ pipeline.

        `lov_file` is the one LOV the job was created with. A job validates
        against that file alone - the upload area is shared, so scanning it
        would pull in reference lists chosen for other jobs.
        """
        required_lovs = self.extract_required_lovs(
            profile_map_path
        )

        available_lovs = self.discover_lov_files(
            lov_dir,
            lov_file=lov_file,
        )

        lov_configuration: dict[str, str] = {}

        for lov_name in required_lovs:
            if lov_name in available_lovs:
                lov_configuration[lov_name] = (
                    available_lovs[lov_name]
                )
                print(
                    f"[INFO] Added required LOV: {lov_name}"
                )
            else:
                print(
                    "[WARN] Required LOV "
                    f"'{lov_name}' not found in uploads/lov/"
                )

        if not required_lovs:
            print(
                "[INFO] No DQ8 rules found, "
                "skipping LOV tables"
            )

        return lov_configuration


lov_service = LovService()
