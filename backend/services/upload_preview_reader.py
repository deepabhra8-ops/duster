from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from utils.logger import get_logger


logger = get_logger(__name__)


class UploadPreviewReader:
    _SUPPORTED_EXTENSIONS = (".csv", ".xlsx")

    def read(
        self,
        file_path: Path,
        rows: int = 50,
    ) -> dict[str, Any]:
        self._validate_rows(rows)

        try:
            self._validate_file(file_path)

            with open(file_path, "rb") as f:
                result = self._read_buffer(f, file_path.suffix.lower(), rows)

        except Exception:
            logger.exception(
                "Failed to read preview file '%s'",
                file_path,
            )
            raise

        logger.debug(
            "Previewed upload '%s' with %d rows",
            file_path,
            len(result["rows"]),
        )

        return result

    @classmethod
    def supports(
        cls,
        file_path: Path,
    ) -> bool:
        suffix = file_path.suffix.lower()

        return suffix in cls._SUPPORTED_EXTENSIONS

    @classmethod
    def supported_extensions(cls) -> tuple[str, ...]:
        return cls._SUPPORTED_EXTENSIONS

    @classmethod
    def _read_buffer(
        cls,
        buffer: Any,
        suffix: str,
        rows: int,
    ) -> dict[str, Any]:
        if suffix not in cls._SUPPORTED_EXTENSIONS:
            supported = ", ".join(cls.supported_extensions())
            raise ValueError(
                f"Unsupported preview file format: {suffix or 'unknown'}. "
                f"Supported formats: {supported}"
            )

        if suffix == ".csv":
            return cls._read_csv_buffer(buffer, rows)

        return cls._read_excel_buffer(buffer, rows)

    @staticmethod
    def _read_csv_buffer(
        buffer: Any,
        rows: int,
    ) -> dict[str, Any]:
        content = buffer.read()
        if isinstance(content, bytes):
            text = content.decode("utf-8", errors="replace")
        else:
            text = content
            
        lines = text.splitlines()
        if not lines:
            return {"columns": [], "rows": []}
            
        reader = csv.reader(lines)
        
        try:
            headers = next(reader)
        except StopIteration:
            return {"columns": [], "rows": []}
            
        columns = [str(h or f"column_{i+1}").strip() for i, h in enumerate(headers)]
        
        preview_rows = []
        for i, row in enumerate(reader):
            if i >= rows:
                break
            preview_rows.append([str(val or "") for val in row])
            
        return {"columns": columns, "rows": preview_rows}

    @staticmethod
    def _read_excel_buffer(
        buffer: Any,
        rows: int,
    ) -> dict[str, Any]:
        workbook = load_workbook(buffer, read_only=True, data_only=True)

        try:
            worksheet = workbook.active
            values = list(
                worksheet.iter_rows(values_only=True, max_row=rows + 1)
            )
        finally:
            workbook.close()

        headers = [
            str(value or "").strip()
            for value in (values[0] if values else ())
        ]
        
        columns = [str(h or f"column_{i+1}").strip() for i, h in enumerate(headers)]
        
        preview_rows = [
            [
                "" if value is None else str(value)
                for value in row
            ]
            for row in values[1:]
        ]

        return {"columns": columns, "rows": preview_rows}

    @staticmethod
    def _validate_rows(
        rows: int,
    ) -> None:
        if not isinstance(rows, int):
            raise ValueError(
                "Preview rows must be an integer"
            )

        if rows < 1:
            raise ValueError(
                "Preview rows must be greater than zero"
            )

    @staticmethod
    def _validate_file(
        file_path: Path,
    ) -> None:
        if not file_path.exists():
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        if not file_path.is_file():
            raise ValueError(
                f"Preview path is not a file: {file_path}"
            )

