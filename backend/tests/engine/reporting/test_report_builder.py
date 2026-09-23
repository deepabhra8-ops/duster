from __future__ import annotations

import pytest
from openpyxl import load_workbook

from engine.core.result_models import (
    DimensionResult,
    RowFailure,
    RuleExecutionDetail,
    TableValidationResult,
    ValidationRunResult,
    ValidationSummary,
)
from engine.reporting.failed_rows_workbook_writer import ROW_ID_COLUMN
from engine.reporting.report_builder import ReportBuilder

pytestmark = pytest.mark.spark


def _detail(rule_id="DQ1", dimension="Completeness", invalid=1, total=4):
    return RuleExecutionDetail(
        table_name="claim",
        column_name="claim_id",
        rule_id=rule_id,
        dimension=dimension,
        cde=True,
        rule_notes="Value must not be null",
        total_rows=total,
        invalid_count=invalid,
        pass_count=total - invalid,
        score=(total - invalid) / total,
    )


@pytest.fixture
def validation_result(spark):
    failed_rows = spark.createDataFrame(
        [("r1", "abc", 10), ("r2", None, 20)],
        [ROW_ID_COLUMN, "claim_id", "amount"],
    )

    table = TableValidationResult(
        table_name="claim",
        total_rows=4,
        pass_mask=None,
        rule_details=[
            _detail(),
            _detail(rule_id="DQ3", dimension="Conformity", invalid=0),
        ],
        row_failures={
            "r2": [RowFailure(column_name="claim_id", rule_id="DQ1", notes="null")],
        },
        failed_rows=failed_rows,
        metadata={"failed_row_count": 2},
    )

    dimensions = {
        "Completeness": DimensionResult.from_scores(dimension="Completeness", scores=[0.75]),
        "Conformity": DimensionResult.from_scores(dimension="Conformity", scores=[1.0]),
    }

    return ValidationRunResult(
        project_name="test-project",
        run_timestamp="2026-01-01T00:00:00",
        summary=ValidationSummary.from_dimensions(dimension_results=dimensions),
        tables={"claim": table},
    )


class TestWorkbookIsProduced:
    def test_the_file_is_written_and_opens(self, validation_result, tmp_path):
        path = ReportBuilder().write(validation_result, tmp_path / "report.xlsx")

        assert path.is_file()
        assert path.stat().st_size > 0

        load_workbook(path)

    def test_it_contains_the_expected_sheets(self, validation_result, tmp_path):
        path = ReportBuilder().write(validation_result, tmp_path / "report.xlsx")

        sheets = load_workbook(path).sheetnames

        assert "Summary" in sheets
        assert any("Completeness" in name for name in sheets)
        assert any("Findings" in name for name in sheets)

    def test_missing_parent_directories_are_created(self, validation_result, tmp_path):
        path = ReportBuilder().write(
            validation_result, tmp_path / "nested" / "deeper" / "report.xlsx"
        )

        assert path.is_file()


class TestContent:
    def test_a_dimension_sheet_carries_its_findings(self, validation_result, tmp_path):
        path = ReportBuilder().write(validation_result, tmp_path / "report.xlsx")
        book = load_workbook(path)

        sheet_name = next(n for n in book.sheetnames if "Completeness" in n)
        values = [
            str(cell.value)
            for row in book[sheet_name].iter_rows()
            for cell in row
            if cell.value is not None
        ]

        assert any("claim_id" in v for v in values)
        assert any("DQ1" in v for v in values)

    def test_no_cell_contains_the_string_nan(self, validation_result, tmp_path):
        path = ReportBuilder().write(validation_result, tmp_path / "report.xlsx")
        book = load_workbook(path)

        for sheet_name in book.sheetnames:
            for row in book[sheet_name].iter_rows():
                for cell in row:
                    assert str(cell.value).lower() != "nan", f"in sheet {sheet_name}"


class TestFormulaInjection:
    def test_a_formula_looking_value_is_stored_as_text(self, spark, tmp_path):
        payload = "=cmd|'/c calc'!A0"

        failed_rows = spark.createDataFrame([("r1", payload)], [ROW_ID_COLUMN, "claim_id"])

        table = TableValidationResult(
            table_name="claim",
            total_rows=1,
            pass_mask=None,
            rule_details=[
                RuleExecutionDetail(
                    table_name="claim",
                    column_name=payload,
                    rule_id="DQ1",
                    dimension="Completeness",
                    cde=False,
                    rule_notes=payload,
                    total_rows=1,
                    invalid_count=1,
                    pass_count=0,
                    score=0.0,
                )
            ],
            row_failures={},
            failed_rows=failed_rows,
            metadata={"failed_row_count": 1},
        )

        result = ValidationRunResult(
            project_name="p",
            run_timestamp="t",
            summary=ValidationSummary.from_dimensions(
                dimension_results={
                    "Completeness": DimensionResult.from_scores(
                        dimension="Completeness", scores=[0.0]
                    )
                }
            ),
            tables={"claim": table},
        )

        path = ReportBuilder().write(result, tmp_path / "report.xlsx")

        book = load_workbook(path)
        formulas = [
            cell.value
            for name in book.sheetnames
            for row in book[name].iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]

        assert formulas == []
