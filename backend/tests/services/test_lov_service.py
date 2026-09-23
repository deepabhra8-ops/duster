from __future__ import annotations

import pytest
from openpyxl import Workbook

from services.lov_service import LovService


RULES_HEADER = "Applicable Rules\n(comma-sep IDs)"
PARAMS_HEADER = "Rule Parameters\n(see Instructions)"


def _write_profile_map(path, rows, sheet_title="Sheet1"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append([RULES_HEADER, PARAMS_HEADER])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_extract_required_lovs_finds_dq8_lov_names(tmp_path):
    profile_map = tmp_path / "profile_map.xlsx"
    _write_profile_map(
        profile_map,
        rows=[
            ["DQ1,DQ8", "CountryList|extra"],
            ["DQ1", ""],
            ["DQ8", "StatusList"],
        ],
    )

    required = LovService().extract_required_lovs(profile_map)

    assert required == {"CountryList", "StatusList"}


def test_extract_required_lovs_skips_the_instructions_sheet(tmp_path):
    profile_map = tmp_path / "profile_map.xlsx"
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Instructions"
    sheet1.append([RULES_HEADER, PARAMS_HEADER])
    sheet1.append(["DQ8", "ShouldBeIgnored"])
    sheet2 = workbook.create_sheet("Data")
    sheet2.append([RULES_HEADER, PARAMS_HEADER])
    sheet2.append(["DQ8", "ShouldBeIncluded"])
    workbook.save(profile_map)

    required = LovService().extract_required_lovs(profile_map)

    assert required == {"ShouldBeIncluded"}


def test_extract_required_lovs_returns_empty_for_a_missing_file(tmp_path):
    missing_path = tmp_path / "does_not_exist.xlsx"
    assert LovService().extract_required_lovs(missing_path) == set()


def test_extract_required_lovs_returns_empty_for_no_path():
    assert LovService().extract_required_lovs(None) == set()


@pytest.mark.spark
def test_discover_lov_files_uses_the_csv_column_headers_as_the_lov_names(tmp_path, spark):
    lov_dir = tmp_path / "lov"
    lov_dir.mkdir()
    (lov_dir / "whatever_filename.csv").write_text("CountryList\nUS\nUK\n")

    discovered = LovService().discover_lov_files(lov_dir)

    assert discovered == {"CountryList": str(lov_dir / "whatever_filename.csv")}


@pytest.mark.spark
def test_discover_lov_files_maps_every_column_of_a_wide_file(tmp_path, spark):
    lov_dir = tmp_path / "lov"
    lov_dir.mkdir()
    lov_file = lov_dir / "job_lovs.csv"
    lov_file.write_text(
        "policy.status,claim.status,customer.gender\n"
        "Active,Open,Male\n"
        "Expired,Closed,Female\n"
        ",Reopened,\n"
    )

    discovered = LovService().discover_lov_files(lov_dir)

    assert discovered == {
        "policy.status": str(lov_file),
        "claim.status": str(lov_file),
        "customer.gender": str(lov_file),
    }


def test_summarize_csv_counts_values_per_column_ignoring_blanks():
    summary = LovService.summarize_csv(
        "policy.status,claim.status\nActive,Open\n,Closed\n,Reopened\n"
    )

    assert summary == [
        {"name": "policy.status", "values_count": 1},
        {"name": "claim.status", "values_count": 3},
    ]


@pytest.mark.spark
def test_discover_lov_files_ignores_non_csv_files(tmp_path, spark):
    lov_dir = tmp_path / "lov"
    lov_dir.mkdir()
    (lov_dir / "notes.txt").write_text("not a csv")

    assert LovService().discover_lov_files(lov_dir) == {}


def test_discover_lov_files_returns_empty_for_a_missing_directory(tmp_path):
    missing_dir = tmp_path / "does_not_exist"
    assert LovService().discover_lov_files(missing_dir) == {}


@pytest.mark.spark
def test_build_lov_configuration_only_includes_lovs_that_are_both_required_and_available(
    tmp_path, spark
):
    profile_map = tmp_path / "profile_map.xlsx"
    _write_profile_map(
        profile_map,
        rows=[["DQ8", "CountryList"], ["DQ8", "MissingList"]],
    )

    lov_dir = tmp_path / "lov"
    lov_dir.mkdir()
    (lov_dir / "countries.csv").write_text("CountryList\nUS\n")

    configuration = LovService().build_lov_configuration(profile_map, lov_dir)

    assert configuration == {"CountryList": str(lov_dir / "countries.csv")}
    assert "MissingList" not in configuration
